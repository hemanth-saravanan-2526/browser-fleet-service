import io
import openpyxl
from engines.extractor import records_to_excel_bytes


def test_records_to_excel_bytes_valid():
    records = [
        {"title": "Product A", "price": "$19.99", "in_stock": True},
        {"title": "Product B", "price": "$29.99", "in_stock": False},
        {"title": "Product C", "price": "$9.99", "in_stock": True},
    ]

    excel_bytes = records_to_excel_bytes(records, sheet_title="Products")
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000

    # Load and verify with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Products"
    assert [cell.value for cell in ws[1]] == ["title", "price", "in_stock"]
    assert ws.max_row == 4
    assert ws.cell(row=2, column=1).value == "Product A"
    assert ws.cell(row=2, column=2).value == "$19.99"


def test_records_to_excel_bytes_single_dict():
    single_record = {"quote": "Simplicity is prerequisite for reliability.", "author": "Dijkstra"}
    excel_bytes = records_to_excel_bytes(single_record, sheet_title="Quotes")
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.max_row == 2
    assert [cell.value for cell in ws[1]] == ["quote", "author"]
