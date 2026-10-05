import io
import json
from typing import Any, Dict, List, Union
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DOM_EXTRACTOR_JS = """
(schema) => {
    function extractValue(root, selector) {
        if (!selector) return "";
        let attr = null;
        let sel = selector;
        if (selector.includes("@")) {
            const parts = selector.split("@");
            sel = parts[0].trim();
            attr = parts[1].trim();
        }

        const elements = sel ? Array.from(root.querySelectorAll(sel)) : [root];
        if (elements.length === 0) return null;

        if (elements.length === 1) {
            const el = elements[0];
            if (attr) {
                return el.getAttribute(attr) || "";
            }
            return (el.innerText || el.textContent || "").trim();
        } else {
            return elements.map(el => {
                if (attr) return el.getAttribute(attr) || "";
                return (el.innerText || el.textContent || "").trim();
            }).filter(v => v !== "");
        }
    }

    const itemSelector = schema._item || schema._item_selector;
    if (itemSelector) {
        const itemElements = Array.from(document.querySelectorAll(itemSelector));
        return itemElements.map(itemEl => {
            const row = {};
            for (const [key, sel] of Object.entries(schema)) {
                if (key.startsWith("_")) continue;
                row[key] = extractValue(itemEl, sel);
            }
            return row;
        });
    } else {
        const result = {};
        for (const [key, sel] of Object.entries(schema)) {
            if (key.startsWith("_")) continue;
            result[key] = extractValue(document, sel);
        }
        return result;
    }
}
"""


def records_to_excel_bytes(data: Union[List[Dict[str, Any]], Dict[str, Any]], sheet_title: str = "Scraped Data") -> bytes:
    """
    Converts structured scraped data into a professionally formatted Excel (.xlsx) spreadsheet.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    # Normalize to list of dictionaries
    if isinstance(data, dict):
        records = [data]
    elif isinstance(data, list):
        records = data
    else:
        records = [{"Value": str(data)}]

    if not records:
        records = [{"Message": "No records extracted"}]

    # Collect all unique headers across all records while preserving order
    headers = []
    for r in records:
        if isinstance(r, dict):
            for k in r.keys():
                if k not in headers:
                    headers.append(k)
        else:
            if "Value" not in headers:
                headers.append("Value")

    # Styling definitions
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")  # Navy Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")  # Light gray
    regular_font = Font(name="Calibri", size=10)
    regular_align = Alignment(vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Write Header Row
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    # Write Data Rows
    for row_idx, r in enumerate(records, start=2):
        row_values = []
        for h in headers:
            val = r.get(h, "") if isinstance(r, dict) else str(r)
            if isinstance(val, list):
                val = ", ".join(str(v) for v in val)
            elif isinstance(val, dict):
                val = json.dumps(val)
            row_values.append(val)

        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 20

        # Style data cells
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.alignment = regular_align
            cell.border = thin_border
            if is_even:
                cell.fill = alt_fill

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            # Don't let single huge cells make columns 500 wide
            line_len = max(len(line) for line in val_str.split("\n")) if val_str else 0
            if line_len > max_len:
                max_len = line_len
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)

    # Save to memory stream
    output_stream = io.BytesIO()
    wb.save(output_stream)
    return output_stream.getvalue()
